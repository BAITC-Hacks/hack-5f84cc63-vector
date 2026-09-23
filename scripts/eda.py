"""Read all source sheets, preserve source files, emit deterministic EDA JSON.

No forecasting, imputation, sign reversal, deduplication or workbook writes.
Roles are structural observations, not business semantics inferred from filenames.
"""
from __future__ import annotations

import argparse
from collections import Counter, defaultdict
from datetime import date, datetime
import hashlib
import itertools
import json
import math
from pathlib import Path
import re
import statistics
import sys

import openpyxl
from openpyxl.utils import get_column_letter

MONTHS = {"янв": 1, "фев": 2, "мар": 3, "апр": 4, "май": 5, "июн": 6,
          "июл": 7, "авг": 8, "сен": 9, "окт": 10, "ноя": 11, "дек": 12}
SKU_HEADERS = {"код", "код 1с", "номенклатура.код"}
TOTAL_RE = re.compile(r"^(?:итого|всего|total|subtotal)(?:\s|:|$)", re.I)


def text(value):
    return "" if value is None else str(value).replace("\xa0", " ").strip()


def numeric(value):
    return isinstance(value, (int, float)) and not isinstance(value, bool) and math.isfinite(value)


def parse_date(value):
    if isinstance(value, datetime):
        return value
    if isinstance(value, date):
        return datetime.combine(value, datetime.min.time())
    if isinstance(value, str):
        for fmt in ("%d.%m.%Y %H:%M:%S", "%d.%m.%Y", "%Y-%m-%d"):
            try:
                return datetime.strptime(value.strip(), fmt)
            except ValueError:
                pass
    return None


def month_label(value):
    match = re.fullmatch(r"([а-яё]+)\.?\s+(20\d{2})", text(value).lower())
    if match and match[1][:3] in MONTHS:
        return f"{match[2]}-{MONTHS[match[1][:3]]:02d}"
    return None


def safe(value):
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    return str(value)


def signature(value):
    return json.dumps(value, ensure_ascii=False, sort_keys=True, default=safe, allow_nan=False)


def digest(path):
    with path.open("rb") as handle:
        return hashlib.file_digest(handle, "sha256").hexdigest()


def quantile(values, q):
    position = (len(values) - 1) * q
    low = int(position)
    high = min(low + 1, len(values) - 1)
    return values[low] + (values[high] - values[low]) * (position - low)


def column_profile(values, rows, label, letter):
    types = Counter()
    nums = []
    errors = []
    dates = []
    for row, value in zip(rows, values):
        if value is None or (isinstance(value, str) and not value.strip()):
            types["null"] += 1
        elif isinstance(value, str) and value.startswith(("#REF!", "#DIV/0!", "#VALUE!", "#N/A", "#NUM!", "#NAME?", "#NULL!")):
            types["excel_error"] += 1
            errors.append({"cell": f"{letter}{row}", "value": value})
        elif isinstance(value, bool):
            types["bool"] += 1
        elif numeric(value):
            types["integer" if float(value).is_integer() else "float"] += 1
            nums.append((row, value))
        elif (dt := parse_date(value)) is not None:
            types["datetime" if isinstance(value, (date, datetime)) else "datetime_string"] += 1
            dates.append(dt.isoformat())
        else:
            types["text"] += 1
    result = {"column": letter, "label": label, "type_counts": dict(types),
              "null_count": types["null"], "rows_profiled": len(values),
              "error_count": len(errors), "error_examples": errors[:5]}
    if dates:
        result["date_range"] = [min(dates), max(dates)]
    if nums:
        ordered = sorted(v for _, v in nums)
        q1, q3 = quantile(ordered, .25), quantile(ordered, .75)
        fence = q3 + 3 * (q3 - q1)
        result["numeric"] = {"count": len(nums), "min": ordered[0], "max": ordered[-1],
            "sum": math.fsum(ordered), "median": statistics.median(ordered),
            "negative": sum(v < 0 for v in ordered), "zero": ordered.count(0),
            "positive": sum(v > 0 for v in ordered),
            "fractional": sum(not float(v).is_integer() for v in ordered),
            "q1": q1, "q3": q3, "upper_3iqr_fence": fence,
            "above_3iqr_count": sum(v > fence for v in ordered),
            "largest_examples": [{"cell": f"{letter}{r}", "value": v} for r, v in sorted(nums, key=lambda p: p[1], reverse=True)[:3]]}
    return result


def identify(values):
    for index, row in enumerate(values[:20]):
        headers = [text(v).lower() for v in row]
        if any(h in SKU_HEADERS for h in headers):
            sku = next(i for i, h in enumerate(headers) if h in SKU_HEADERS)
            months = {i: month_label(v) for i, v in enumerate(row) if month_label(v)}
            if "дата" in headers and "количество" in headers:
                role = "transactions"
            elif any("поступление до" in h for h in headers):
                role = "dated_shipments"
            elif months:
                subheaders = " ".join(text(v).lower() for r in values[index + 1:index + 3] for v in r)
                role = "monthly_beginning_inventory" if "нач. остаток" in subheaders else "monthly_quantity_unspecified"
            elif "кратность" in headers or any("мин. разр." in h for h in headers):
                role = "order_constraints"
            else:
                role = "sku_table_unclassified"
            return index, sku, months, role
    return None, None, {}, "aggregate_seasonality" if any("коэф. сезонности" in text(v).lower() for row in values for v in row) else "unclassified"


def analyze_sheet(ws, formula_ws):
    values = list(ws.iter_rows(values_only=True))
    formula_count = 0
    missing_caches = []
    external_formulas = []
    formulas_sample = []
    raw_errors = []
    empty_error_typed = []
    header_index, sku_col, months, role = identify(values)
    for r, row in enumerate(formula_ws.iter_rows(), 1):
        for c, cell in enumerate(row):
            cached = values[r - 1][c]
            if cell.data_type == "f":
                formula_count += 1
                item = {"cell": cell.coordinate, "formula": cell.value, "cached": safe(cached)}
                if len(formulas_sample) < 8:
                    formulas_sample.append(item)
                if cached is None:
                    missing_caches.append(cell.coordinate)
                if "[" in str(cell.value) and "]" in str(cell.value):
                    external_formulas.append(item)
            if cell.data_type == "e" and cached is None:
                empty_error_typed.append(cell.coordinate)
            elif cell.data_type == "e" or (isinstance(cached, str) and cached.startswith(("#N/A", "#DIV/0!", "#REF!", "#VALUE!", "#NUM!", "#NAME?"))):
                raw_errors.append({"cell": cell.coordinate, "value": safe(cached)})
    nonempty = [(r, row) for r, row in enumerate(values, 1) if any(v is not None for v in row)]
    result = {"sheet": ws.title, "state": ws.sheet_state, "dimensions": [ws.max_row, ws.max_column],
        "observed_nonempty_rows": len(nonempty), "blank_rows": len(values) - len(nonempty),
        "role_from_structure": role, "header_row": header_index + 1 if header_index is not None else None,
        "formulas": {"count": formula_count, "examples": formulas_sample,
                     "missing_cache_count": len(missing_caches), "missing_cache_examples": missing_caches[:10],
                     "external_reference_count": len(external_formulas), "external_reference_examples": external_formulas[:5]},
        "excel_errors": {"count": len(raw_errors), "examples": raw_errors[:10]},
        "empty_error_typed_cells": {"count": len(empty_error_typed), "examples": empty_error_typed[:10],
            "meaning": "XML type=e without value; counted as null, not as a displayed Excel error"},
        "header_evidence": [{"row": r, "cells": {get_column_letter(c + 1): safe(v) for c, v in enumerate(row) if v is not None}} for r, row in nonempty[:3]],
        "content_sha256": hashlib.sha256(signature(values).encode()).hexdigest()}
    if header_index is None:
        # Layouts with multiple small tables: whole-sheet profiles are deliberately
        # not interpreted as one transaction table. Retain exact cells for audit.
        result["profile_scope"] = "whole_sheet_multi_table; duplicate_rows_not_business_duplicates"
        result["columns"] = [column_profile([row[c] for row in values], list(range(1, len(values) + 1)), "multiple labels", get_column_letter(c + 1)) for c in range(ws.max_column)]
        result["cells"] = [{"row": r, "cells": {get_column_letter(c + 1): safe(v) for c, v in enumerate(row) if v is not None}} for r, row in nonempty]
        result["duplicate_nonempty_rows"] = len(nonempty) - len({signature(row) for _, row in nonempty})
        result["sku_count"] = None
        result["date_range"] = None
        result["years_in_column_A"] = sorted({int(row[0]) for _, row in nonempty if numeric(row[0]) and 2000 <= row[0] <= 2100})
        result["total_label_cells"] = [f"{get_column_letter(c + 1)}{r}" for r, row in nonempty for c, v in enumerate(row) if isinstance(v, str) and TOTAL_RE.match(v.strip())]
        return result, None
    headers = [text(v) or f"unnamed_{i+1}" for i, v in enumerate(values[header_index])]
    included, totals, excluded = [], [], []
    for r, row in nonempty:
        if r <= header_index + 1:
            continue
        # Whole-word start match, not substring inside an arbitrary product name.
        markers = [get_column_letter(c + 1) for c, v in enumerate(row) if isinstance(v, str) and TOTAL_RE.match(v.strip())]
        if markers:
            totals.append({"row": r, "marker_columns": markers, "values": [safe(v) for v in row]})
        elif not text(row[sku_col]):
            excluded.append({"row": r, "reason": "no_sku; may_be_subheader_or_artifact"})
        else:
            included.append((r, row))
    result.update({"profile_scope": "rows_with_nonblank_sku_excluding_total_markers; duplicates_retained",
        "columns": [column_profile([row[c] for _, row in included], [r for r, _ in included], h, get_column_letter(c+1)) for c,h in enumerate(headers)],
        "data_rows": len(included), "total_subtotal_rows": totals, "other_excluded_rows": excluded,
        "blank_rows_below_header": sum(all(v is None for v in row) for row in values[header_index+1:]),
        "month_columns": {get_column_letter(c+1): m for c,m in months.items()},
        "month_range": [min(months.values()), max(months.values())] if months else None})
    sku_counter = Counter(text(row[sku_col]) for _, row in included)
    result["sku_count"] = len(sku_counter)
    result["join_keys"] = {"sku_column": get_column_letter(sku_col+1), "sku_header": headers[sku_col],
        "normalization": "trim whitespace only; preserve leading zeros and underscores",
        "distinct_codes": sorted(sku_counter), "repeated_code_count": sum(n>1 for n in sku_counter.values()),
        "alternative_headers": [h for h in headers if "артикул" in h.lower()],
        "numeric_sku_cells": sum(numeric(row[sku_col]) for _,row in included)}
    stripped = defaultdict(set)
    for sku in sku_counter:
        stripped[sku.rstrip("_")].add(sku)
    result["join_keys"]["trailing_underscore_removal_collisions"] = {k: sorted(v) for k,v in stripped.items() if len(v)>1}
    seen = {}
    duplicates = []
    for r,row in included:
        key = signature(row)
        if key in seen:
            duplicates.append({"row": r, "first_row": seen[key]})
        else:
            seen[key] = r
    result["exact_duplicate_rows"] = {"count": len(duplicates), "examples": duplicates[:10]}
    pack = {"skus": set(sku_counter), "months": {}, "constraints": {}, "role": role, "result": result}
    constraint_col = next((i for i,h in enumerate(headers) if h.lower() == "кратность" or "мин. разр." in h.lower()), None)
    if constraint_col is not None:
        constraint_rows = defaultdict(list)
        for r,row in included:
            constraint_rows[text(row[sku_col])].append((r,row[constraint_col]))
        conflicts = []
        for sku, entries in constraint_rows.items():
            distinct = {signature(v) for _,v in entries}
            pack["constraints"][sku] = entries[0][1] if len(distinct)==1 else None
            if len(distinct)>1:
                conflicts.append({"sku":sku,"rows_and_values":entries})
        result["conflicting_constraint_codes"] = conflicts
        cv = list(pack["constraints"].values())
        result["constraint_values"] = {"label": headers[constraint_col], "positive_numeric_skus": sum(numeric(v) and v>0 for v in cv),
            "zero_skus": sum(numeric(v) and v==0 for v in cv), "negative_skus": sum(numeric(v) and v<0 for v in cv),
            "missing_skus": sum(v is None for v in cv), "nonnumeric_skus": sum(v is not None and not numeric(v) for v in cv)}
    result["coefficient_headers"] = [h for h in headers if re.search("рост|сез|коэф|кэф", h, re.I)]
    result["business_field_headers"] = [h for h in headers if re.search("клиент|цен[аы]|категор|групп|резерв|свобод|срок|путь|витрин|остат", h, re.I)]
    if role == "transactions":
        dc, qc = headers.index("Дата"), headers.index("Количество")
        signs = Counter()
        years = defaultdict(Counter)
        dates = []
        invalid = []
        doc_keys = Counter()
        for r,row in included:
            dt, qty = parse_date(row[dc]), row[qc]
            if dt is None or not numeric(qty):
                invalid.append(r)
                continue
            dates.append(dt.isoformat())
            sign = "positive" if qty>0 else "negative" if qty<0 else "zero"
            signs[sign] += 1
            years[str(dt.year)][sign] += 1
            key = (text(row[sku_col]), dt.strftime("%Y-%m"))
            pack["months"][key] = pack["months"].get(key, 0) + qty
            doc_keys[tuple(text(row[headers.index(h)]) for h in ("Дата", "Номер", "Код", "Склад"))] += 1
        result["transactions"] = {"sign_counts": dict(signs), "sign_by_year": {y:dict(v) for y,v in sorted(years.items())},
            "invalid_date_or_qty_rows": invalid, "date_range": [min(dates),max(dates)] if dates else None,
            "repeated_date_document_sku_warehouse_keys": sum(v>1 for v in doc_keys.values()),
            "warehouse_counts": dict(Counter(text(row[headers.index("Склад")]) for _,row in included)),
            "unit_counts": dict(Counter(text(row[headers.index("Ед.")]) for _,row in included))}
        result["date_range"] = result["transactions"]["date_range"]
    elif months:
        result["date_range"] = None
        for _,row in included:
            for c,m in months.items():
                if numeric(row[c]):
                    key = (text(row[sku_col]),m)
                    pack["months"][key] = pack["months"].get(key,0) + row[c]
        total_col = next((i for i,h in enumerate(headers) if h.lower()=="итого"),None)
        if total_col is not None:
            first_c = min(months)
            eligible = [(r,row) for r,row in included if numeric(row[total_col])]
            result["total_column_checks"] = {"numeric_total_rows": len(eligible),
                "equal_first_month_with_both_numeric": sum(numeric(row[first_c]) and math.isclose(row[total_col],row[first_c],rel_tol=0,abs_tol=1e-6) for _,row in eligible),
                "both_total_first_month_numeric": sum(numeric(row[first_c]) for _,row in eligible),
                "equal_sum_of_available_numeric_months": sum(math.isclose(row[total_col],math.fsum(row[c] for c in months if numeric(row[c])),rel_tol=0,abs_tol=1e-6) for _,row in eligible),
                "caution": "available-cell sum diagnostic only; blanks not imputed"}
    if role == "dated_shipments":
        shipment_columns = []
        for c,h in enumerate(headers):
            if "поступление до" in h.lower():
                match = re.search(r"поступление до\s+(\d{2}\.\d{2}\.\d{4})",h,re.I)
                dt = parse_date(match[1]) if match else None
                ids = re.findall(r"УТ-\d+",h)
                shipment_columns.append({"column": get_column_letter(c+1),"header": h,"shipment_id": ids[0] if ids else None,
                    "arrival_deadline": dt.date().isoformat() if dt else None,
                    "positive_rows": sum(numeric(row[c]) and row[c]>0 for _,row in included)})
        result["shipment_columns"] = shipment_columns
        result["date_range"] = None
    return result, pack


def cross_checks(packs):
    intersections, coverage, reconciliations, conflicts = [], [], [], []
    for a,b in itertools.combinations(packs,2):
        if a["supplier"] != b["supplier"]:
            continue
        common = a["skus"] & b["skus"]
        intersections.append({"a":a["id"],"b":b["id"],"shared_skus":len(common),
            "only_a":len(a["skus"]-b["skus"]),"only_b":len(b["skus"]-a["skus"])})
        if a["constraints"] and b["constraints"]:
            differences=[{"sku":k,"a":safe(a["constraints"][k]),"b":safe(b["constraints"][k])} for k in sorted(common) if a["constraints"][k] != b["constraints"][k]]
            conflicts.append({"a":a["id"],"b":b["id"],"shared_skus":len(common),"different_values":len(differences),"examples":differences[:10]})
        if (a["role"]=="transactions") != (b["role"]=="transactions"):
            tx,other=(a,b) if a["role"]=="transactions" else (b,a)
            if other["months"]:
                keys=sorted(tx["months"].keys() & other["months"].keys())
                diffs=[{"sku":k[0],"month":k[1],"transactions_signed":tx["months"][k],"monthly_value":other["months"][k]} for k in keys if not math.isclose(tx["months"][k],other["months"][k],rel_tol=0,abs_tol=1e-6)]
                reconciliations.append({"transactions":tx["id"],"monthly":other["id"],"numeric_pairs_compared":len(keys),"equal_pairs":len(keys)-len(diffs),"different_pairs":len(diffs),
                    "examples":diffs[:10],"method":"signed transactions; duplicates retained; only pairs with numeric monthly cell and observed transactions; no blank=zero assumption"})
    for source in packs:
        if source["role"] != "order_constraints":
            continue
        positive={k for k,v in source["constraints"].items() if numeric(v) and v>0}
        for target in packs:
            if target["supplier"] != source["supplier"] or target is source:
                continue
            absent=target["skus"]-source["skus"]
            invalid=(target["skus"] & source["skus"])-positive
            coverage.append({"constraint_source":source["id"],"target":target["id"],"target_skus":len(target["skus"]),
                "matched_skus":len(target["skus"] & source["skus"]),"absent_skus":len(absent),"invalid_or_missing_value_skus":len(invalid),
                "usable_positive_skus":len(target["skus"] & positive),"absent_codes":sorted(absent),"invalid_codes":sorted(invalid)})
    return {"sku_intersections":intersections,"constraint_coverage":coverage,"constraint_conflicts":conflicts,"monthly_transaction_comparisons":reconciliations}


def markdown_report(report):
    lines=["# Verified Excel data map", "", "Generated by `scripts/eda.py` from every sheet. Sources are opened read-only.",
        "Filenames do not establish table semantics. `monthly_quantity_unspecified` requires business clarification.",
        "Codes are joined after trimming whitespace, preserving `_` and leading zeros.",
        "", f"Files: {len(report['files'])}. JSON: `data/processed/eda_summary.json` contains SHA-256, column profiles, errors, duplicates, and coverage.",
        "Formulas were not recalculated: cached values, errors, and missing caches were checked separately.",
        "Numeric 3×IQR outliers are investigation candidates, not proven one-off orders.", ""]
    for f in report["files"]:
        lines += [f"## {f['path']}", "", f"SHA-256: `{f['sha256']}`", ""]
        for s in f["sheets"]:
            lines += [f"### Sheet `{s['sheet']}`", "", f"Dimensions: {s['dimensions'][0]} × {s['dimensions'][1]}; type: `{s['role_from_structure']}`.",
                f"Nonempty rows: {s['observed_nonempty_rows']}; SKUs: {s['sku_count'] if s['sku_count'] is not None else 'not applicable'}; Excel errors: {s['excel_errors']['count']}."]
            if s.get("header_row"):
                lines += [f"Header row: {s['header_row']}; SKU rows excluding totals: {s['data_rows']}; totals: {len(s['total_subtotal_rows'])}; exact duplicate rows: {s['exact_duplicate_rows']['count']}."]
            if s.get("month_range"):
                lines += [f"Monthly columns: {' — '.join(s['month_range'])}."]
            if s.get("date_range"):
                lines += [f"Transaction dates: {' — '.join(s['date_range'])}."]
            if s.get("transactions"):
                lines += [f"Quantity signs: `{json.dumps(s['transactions']['sign_counts'])}`."]
            lines += ["", "| Column | Source header | Types (cell counts) | Nulls |", "|---|---|---|---|"]
            for c in s["columns"]:
                label=c["label"].replace("|","/").replace("\n"," ")
                types=", ".join(f"{k}: {v}" for k,v in c["type_counts"].items() if v)
                label = f'"{label}"' if label != "multiple labels" else label
                lines.append(f"| {c['column']} | {label} | {types} | {c['null_count']} |")
            lines += [""]
    lines += ["## Order constraint coverage", "", "Coverage checks both code presence and a positive numeric value.", ""]
    for c in report["cross_checks"]["constraint_coverage"]:
        lines.append(f"- `{c['constraint_source']}` → `{c['target']}`: {c['usable_positive_skus']}/{c['target_skus']} with positive values; absent codes {c['absent_skus']}; present codes with invalid values {c['invalid_or_missing_value_skus']}.")
    return "\n".join(lines)+"\n"


def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--data-dir",type=Path,default=Path("data"))
    parser.add_argument("--output",type=Path,default=Path("data/processed/eda_summary.json"))
    parser.add_argument("--markdown",type=Path,default=Path("docs/DATA.md"))
    args=parser.parse_args()
    paths=sorted(p for p in args.data_dir.rglob("*.xlsx") if "processed" not in p.relative_to(args.data_dir).parts and not p.name.startswith("~$"))
    if not paths:
        parser.error("No input xlsx files found")
    report={"report_version":1,"python_version":sys.version.split()[0],"openpyxl_version":openpyxl.__version__,
        "method":"read-only, all sheets and rows; cached formulas inspected separately; no business imputation",
        "limitations":["3IQR per numeric column mixes SKUs/units; diagnostic only, never automatic outlier removal",
                       "aggregate multi-table sheets have physical-column profiles, not semantic business-column dtypes",
                       "numeric pairs only in monthly comparisons; scope and blanks must be clarified before reconciliation",
                       "source formula caches are read, not recalculated"], "files":[]}
    packs=[]
    for path in paths:
        before=digest(path)
        print(f"Analyzing {path}",flush=True)
        supplier=path.relative_to(args.data_dir).parts[0]
        cached=openpyxl.load_workbook(path,read_only=True,data_only=True,keep_links=False)
        formulas=openpyxl.load_workbook(path,read_only=True,data_only=False,keep_links=False)
        f={"path":path.as_posix(),"supplier_folder":supplier,"sha256":before,"bytes":path.stat().st_size,"sheets":[]}
        try:
            for ws in cached:
                s,pack=analyze_sheet(ws,formulas[ws.title])
                f["sheets"].append(s)
                if pack:
                    pack.update({"id":f"{path.as_posix()}::{ws.title}","supplier":supplier})
                    packs.append(pack)
        finally:
            cached.close()
            formulas.close()
        f["source_hash_unchanged"]=before==digest(path)
        if not f["source_hash_unchanged"]:
            raise RuntimeError(f"Source changed during analysis: {path}")
        report["files"].append(f)
    report["cross_checks"]=cross_checks(packs)
    by_hash=defaultdict(list)
    by_sheet=defaultdict(list)
    for f in report["files"]:
        by_hash[f["sha256"]].append(f["path"])
        for s in f["sheets"]:
            by_sheet[s["content_sha256"]].append(f"{f['path']}::{s['sheet']}")
    report["identical_files"]=[v for v in by_hash.values() if len(v)>1]
    report["identical_cached_sheets"]=[v for v in by_sheet.values() if len(v)>1]
    args.output.parent.mkdir(parents=True,exist_ok=True)
    args.output.write_text(json.dumps(report,ensure_ascii=False,indent=2,default=safe,allow_nan=False)+"\n",encoding="utf-8")
    args.markdown.parent.mkdir(parents=True,exist_ok=True)
    args.markdown.write_text(markdown_report(report),encoding="utf-8")
    print(f"Complete: {len(paths)} files -> {args.output}",flush=True)


if __name__ == "__main__":
    main()
