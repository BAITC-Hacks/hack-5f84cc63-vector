"""Stream canonical JSONL tables into a new, independently auditable run."""
from collections import Counter, defaultdict
from datetime import datetime, timezone
import hashlib
import json
from pathlib import Path
from uuid import uuid4

from openpyxl import load_workbook

from .adapters import IEKAdapter, SystemElectricAdapter, SourceLayoutError
from .config import load_config, sha256
from .schema import SCHEMAS, validate


def dump(path, value):
    Path(path).write_text(json.dumps(value,ensure_ascii=False,indent=2,allow_nan=False)+"\n",encoding="utf-8")


class CanonicalWriter:
    def __init__(self, directory):
        self.directory=directory
        self.handles={table:(directory/f"{table}.jsonl").open("x",encoding="utf-8",newline="\n") for table in SCHEMAS}
        self.counts=Counter()
        self.flags=defaultdict(Counter)
        self.sources=defaultdict(Counter)
        self.product_facts={}
        self.constraint_groups=defaultdict(list)
        self.transaction_stats=defaultdict(Counter)

    def write(self,table,record):
        validate(table,record)
        self.handles[table].write(json.dumps(record,ensure_ascii=False,separators=(",",":"),allow_nan=False)+"\n")
        self.counts[table]+=1
        self.flags[table].update(record["quality_flags"])
        if "source_id" in record:
            self.sources[record["source_id"]][table]+=1
        if table=="product_observations":
            key=(record["supplier_id"],record["sku_id"])
            facts=self.product_facts.setdefault(key,{"raw":set(),"fields":defaultdict(set),"refs":[]})
            facts["raw"].add(json.dumps(record["sku_raw"],ensure_ascii=False))
            for field in ("name","supplier_article","stock_uom"):
                if record[field] is not None:
                    facts["fields"][field].add(record[field])
            facts["refs"].append({"record_id":record["record_id"],"source_id":record["source_id"],"source_row":record["source_row"]})
        if table=="order_constraints":
            self.constraint_groups[(record["supplier_id"],record["sku_id"])].append({
                "source_id":record["source_id"],"source_row":record["source_row"],
                "value":record["constraint_value_raw"],"value_status":record["value_status"],"constraint_kind":record["constraint_kind"]})
        if table=="transactions":
            stats=self.transaction_stats[record["supplier_id"]]
            stats["rows"]+=1
            qty=record["qty_raw"]
            stats["missing_quantity" if qty is None else "negative" if qty<0 else "positive" if qty>0 else "zero"]+=1

    def finish_products(self):
        for (supplier,sku),facts in sorted(self.product_facts.items()):
            candidates={k:sorted(facts["fields"][k]) for k in ("name","supplier_article","stock_uom")}
            constraints=self.constraint_groups.get((supplier,sku),[])
            flags=["category_unknown"]
            for key,values in candidates.items():
                if len(values)>1:
                    flags.append(f"conflicting_{key}")
                elif not values:
                    flags.append(f"missing_{key}")
            if not constraints:
                flags.append("constraint_missing")
            elif not any(c["value_status"]=="valid" for c in constraints):
                flags.append("constraint_unusable")
            if len({json.dumps(c["value"],ensure_ascii=False) for c in constraints})>1:
                flags.append("constraint_source_conflict")
            self.write("products",{"supplier_id":supplier,"sku_id":sku,
                "sku_raw_variants":[json.loads(v) for v in sorted(facts["raw"])],
                **{k:v[0] if len(v)==1 else None for k,v in candidates.items()},
                "category_id":None,"attribute_candidates":candidates,"source_refs":facts["refs"],
                "quality_flags":flags,"is_synthetic":False})

    def close(self):
        for handle in self.handles.values():
            handle.close()


def run(manifest_path="config/sources.json",settings_path="config/ingestion.json",output_dir=None,root="."):
    root=Path(root).resolve()
    manifest_path=(root/manifest_path).resolve()
    settings_path=(root/settings_path).resolve()
    manifest,settings=load_config(manifest_path,settings_path)
    config_hashes={"manifest":sha256(manifest_path),"settings":sha256(settings_path)}
    by_id={s["source_id"]:s for s in manifest["sources"]}
    paths={}
    active_identities=set()
    for source in manifest["sources"]:
        path=(root/source["path"]).resolve()
        if not path.is_relative_to(root/"data") or "processed" in path.relative_to(root/"data").parts:
            raise ValueError("Sources must be inside the raw data tree")
        if path not in paths:
            paths[path]=sha256(path)
        if paths[path]!=source["sha256"]:
            raise SourceLayoutError(f"Source fingerprint changed; review mapping first: {source['source_id']}")
        if source.get("duplicate_of"):
            original=by_id[source["duplicate_of"]]
            if original.get("duplicate_of") or original["sha256"]!=source["sha256"] or original["sheet"]!=source["sheet"] or original["supplier_id"]!=source["supplier_id"]:
                raise ValueError("Duplicate source must refer directly to an identical source sheet")
        else:
            identity=(source["supplier_id"],source["sha256"],source["sheet"])
            if identity in active_identities:
                raise ValueError("Identical source sheets must declare duplicate_of")
            active_identities.add(identity)
    discovered={p.resolve() for p in (root/"data").rglob("*.xlsx") if "processed" not in p.relative_to(root/"data").parts and not p.name.startswith("~$")}
    if discovered != set(paths):
        raise SourceLayoutError("Source registry does not match raw files; review new or missing sources")
    started=datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%SZ")
    directory=(root/output_dir).resolve() if output_dir else root/"data"/"processed"/"canonical"/f"{started}-{uuid4().hex[:8]}"
    allowed=root/"data"/"processed"
    if not directory.is_relative_to(allowed) or directory==allowed:
        raise ValueError("Canonical run directory must be a new child of data/processed")
    directory.mkdir(parents=True,exist_ok=False)
    writer=CanonicalWriter(directory)
    adapters={"IEK":IEKAdapter(settings),"SystemElectric":SystemElectricAdapter(settings)}
    source_status=[]
    try:
        for source in manifest["sources"]:
            if source.get("duplicate_of"):
                source_status.append({"source_id":source["source_id"],"status":"skipped_identical_duplicate","duplicate_of":source["duplicate_of"]})
                continue
            print(f"Ingesting {source['source_id']}",flush=True)
            book=load_workbook(root/source["path"],read_only=True,data_only=True,keep_links=False)
            try:
                if source["sheet"] not in book.sheetnames:
                    raise SourceLayoutError(f"Missing source sheet: {source['source_id']}")
                for table,record in adapters[source["supplier_id"]].parse(source,book[source["sheet"]].iter_rows(values_only=True)):
                    writer.write(table,record)
            finally:
                book.close()
            source_status.append({"source_id":source["source_id"],"status":"processed"})
        writer.finish_products()
        writer.close()
        # Do not publish a complete report for a changing input snapshot.
        for path,expected in paths.items():
            if sha256(path)!=expected:
                raise RuntimeError(f"Source changed during ingestion: {path.name}")
        if sha256(manifest_path)!=config_hashes["manifest"] or sha256(settings_path)!=config_hashes["settings"]:
            raise RuntimeError("Configuration changed during ingestion")
        conflict_groups=[{"supplier_id":supplier,"sku_id":sku,"values":values} for (supplier,sku),values in sorted(writer.constraint_groups.items())
            if len({json.dumps(v["value"],ensure_ascii=False) for v in values})>1]
        artifacts={f"{table}.jsonl":{"rows":writer.counts[table],"sha256":sha256(directory/f"{table}.jsonl")} for table in SCHEMAS}
        report={"status":"complete","schema_version":1,"started_at_utc":started,
            "config_sha256":config_hashes,"source_files":len(paths),"source_sheets":len(manifest["sources"]),
            "sources":source_status,"tables":artifacts,"rows_by_source":dict(writer.sources),
            "quality_flag_counts":dict(writer.flags),"transaction_stats":dict(writer.transaction_stats),
            "constraint_conflicts":conflict_groups,
            "duplicate_groups":[g for adapter in adapters.values() for g in adapter.duplicate_groups],
            "row_audit":{supplier:dict(adapter.audit) for supplier,adapter in adapters.items()},
            "source_hashes_unchanged":True,"effective_settings":settings,
            "limitations":["No order quantities are calculated by ingestion.",
                "Missing current stock, SE shipments, categories, customer IDs, and unresolved supplier policies remain missing.",
                "Alternative monthly sources and coefficient series remain separate; no automatic precedence or aggregation.",
                "Inspect quality flags and conflict groups before downstream calculations.",
                "Source formula caches are read without recalculation."]}
        dump(directory/"source_manifest.json",manifest)
        dump(directory/"schema.json",SCHEMAS)
        dump(directory/"ingestion_summary.json",report)
        return directory,report
    except Exception as exc:
        writer.close()
        dump(directory/"FAILED.json",{"status":"failed","error":str(exc)})
        raise
